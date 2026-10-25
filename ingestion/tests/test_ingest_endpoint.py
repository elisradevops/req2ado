# ingestion/tests/test_ingest_endpoint.py
import io
import tempfile
from pathlib import Path

from fastapi.testclient import TestClient
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from app.config import settings
from app.main import _figure_caption_cache_key, _rasterize_for_vision, _sniff_image_mime_type, app
from app.vision.orchestrator_vision_client import PROMPT_VERSION, FigureVisionResult

from tests.identity import identity_headers

client = TestClient(app)
FIXTURES = Path(__file__).parent / "fixtures"


def _upload(filename: str, path: Path, redact: bool = False, **extra_params: object) -> dict:
    with open(path, "rb") as f:
        response = client.post(
            "/ingest",
            # user_id defaults on: vision only runs for a request that carries
            # the submitting user's identity (the orchestrator's /llm/vision
            # needs it to resolve the user's credential/model).
            params=(params := {"redact": redact, "user_id": "test-user", **extra_params}),
            headers=identity_headers(params),
            files={"file": (filename, f, "application/octet-stream")},
        )
    assert response.status_code == 200
    return response.json()


class _FakeRedis:
    # Module-level shared fake -- same minimal get/set/publish shape as the
    # per-test local _FakeRedis classes elsewhere in this file, reused by
    # tests that only need a working cache, not narration-publish assertions.
    def __init__(self) -> None:
        self._store: dict[str, bytes] = {}

    def get(self, key):
        return self._store.get(key)

    def set(self, key, value, ex=None):
        self._store[key] = value.encode("utf-8") if isinstance(value, str) else value

    def publish(self, channel, payload):
        pass


def test_ingest_txt() -> None:
    body = _upload("sample.txt", FIXTURES / "sample.txt")
    assert body["document"]["format"] == "txt"
    assert len(body["blocks"]) == 2


def test_ingest_markdown() -> None:
    body = _upload("sample.md", FIXTURES / "sample.md")
    assert body["document"]["format"] == "md"
    assert any(b["kind"] == "heading" for b in body["blocks"])


def test_ingest_docx() -> None:
    body = _upload("sample.docx", FIXTURES / "sample.docx")
    assert body["document"]["format"] == "docx"
    assert len(body["tables"]) == 1


def test_ingest_xlsx() -> None:
    body = _upload("sample.xlsx", FIXTURES / "sample.xlsx")
    assert body["document"]["format"] == "xlsx"
    assert len(body["tables"]) == 2


def test_ingest_pdf() -> None:
    body = _upload("sample.pdf", FIXTURES / "sample.pdf")
    assert body["document"]["format"] == "pdf"
    assert len(body["blocks"]) >= 1


def test_ingest_scanned_pdf_via_ocr() -> None:
    body = _upload("scanned_sample.pdf", FIXTURES / "scanned_sample.pdf")
    assert body["document"]["format"] == "pdf_scanned"
    assert any(b["kind"] == "open_tbd" for b in body["blocks"])


def test_ingest_pptx() -> None:
    body = _upload("sample.pptx", FIXTURES / "sample.pptx")
    assert body["document"]["format"] == "pptx"
    assert any(b["kind"] == "heading" for b in body["blocks"])


def test_ingest_csv() -> None:
    body = _upload("sample.csv", FIXTURES / "sample.csv")
    assert body["document"]["format"] == "csv"
    assert len(body["tables"]) == 1
    assert [cell["text"] for cell in body["tables"][0]["cells"][0]] == ["Component", "Voltage", "Current (A)"]


def test_ingest_rtf() -> None:
    body = _upload("sample.rtf", FIXTURES / "sample.rtf")
    assert body["document"]["format"] == "rtf"
    assert any("thermal requirements" in b["text"] for b in body["blocks"])


def test_ingest_image() -> None:
    body = _upload("sample.png", FIXTURES / "sample.png")
    assert body["document"]["format"] == "image"
    assert len(body["blocks"]) == 1
    assert body["blocks"][0]["kind"] == "figure"


def test_ingest_with_redact_strips_emails() -> None:
    # sample.txt has no email, so use a temp fixture with one inline
    import tempfile

    with tempfile.NamedTemporaryFile(suffix=".txt", mode="w", delete=False) as f:
        f.write("Contact jane.doe@example.com for details.")
        temp_path = Path(f.name)

    body = _upload(temp_path.name, temp_path, redact=True)
    assert "jane.doe@example.com" not in body["blocks"][0]["text"]
    temp_path.unlink()


def test_ingest_rejects_a_file_exceeding_the_configured_size_cap(monkeypatch) -> None:
    # Regression: /ingest previously buffered the whole upload into memory
    # with no size guard of its own at this layer (the orchestrator's own
    # check happens before forwarding here, but this endpoint is
    # independently reachable). Set a tiny cap so a small fixture trips it
    # without needing to actually construct a huge file.
    monkeypatch.setattr(settings, "max_upload_size_mb", 0)  # effectively 0 bytes allowed
    with tempfile.NamedTemporaryFile(suffix=".txt", mode="w", delete=False) as f:
        f.write("more than zero bytes of content")
        temp_path = Path(f.name)

    with open(temp_path, "rb") as f:
        response = client.post(
            "/ingest", files={"file": (temp_path.name, f, "application/octet-stream")}
        )
    assert response.status_code == 413
    temp_path.unlink()


def test_ingest_skip_vision_skips_figure_captioning_even_when_vision_is_configured(monkeypatch) -> None:
    # Regression: /jobs/preview (orchestrator) only needs blockCount/
    # sectionCount/fileType -- it never reads figure captions -- but before
    # skip_vision existed, every embedded image in a merely-previewed
    # document paid the full per-image vision-LLM inference cost anyway
    # (confirmed live: a single small image made preview take 28s instead
    # of ~30ms). skip_vision=true must suppress describe_figure entirely,
    # even when a real vision client IS configured.
    import app.main as main_module

    class _FakeVisionClient:
        def __init__(self) -> None:
            self.calls = 0

        def describe_figure(self, image_bytes: bytes, mime_type: str | None = None, user_id: str | None = None) -> FigureVisionResult | None:
            self.calls += 1
            return FigureVisionResult(significant=True, description="a description nobody asked for")

    fake_client = _FakeVisionClient()
    monkeypatch.setattr(main_module, "_vision_client", fake_client)

    body = _upload("sample.docx", FIXTURES / "sample.docx", **{"skip_vision": True})
    figure_blocks = [b for b in body["blocks"] if b["kind"] == "figure"]
    assert len(figure_blocks) == 1
    assert figure_blocks[0]["text"] == "[FIGURE: none]"  # placeholder, not the fake caption
    assert fake_client.calls == 0


def test_ingest_survey_mode_still_captions_when_vision_is_configured(monkeypatch) -> None:
    # Companion to the test above -- proves skip_vision is opt-in, not a
    # silent global regression: real SURVEY job submission
    # (pipeline.worker.ts, mode='survey') must keep getting real captions.
    import app.main as main_module

    class _FakeVisionClient:
        def __init__(self) -> None:
            self.calls = 0

        def describe_figure(self, image_bytes: bytes, mime_type: str | None = None, document_context: str | None = None, user_id: str | None = None) -> FigureVisionResult | None:
            self.calls += 1
            return FigureVisionResult(significant=True, description="a real caption")

    fake_client = _FakeVisionClient()
    monkeypatch.setattr(main_module, "_vision_client", fake_client)

    body = _upload("sample.docx", FIXTURES / "sample.docx", mode="survey")
    figure_blocks = [b for b in body["blocks"] if b["kind"] == "figure"]
    assert figure_blocks[0]["text"] == "[FIGURE: a real caption]"
    assert fake_client.calls == 1


def test_describe_figure_caches_by_image_hash_so_a_rerun_gets_the_same_caption(monkeypatch) -> None:
    # Regression guard for a confirmed real cause of run-to-run SURVEY
    # node-count variance: a figure's caption used to be freshly (re-)
    # sampled on every job, and caption text feeds llm_leveling.py's
    # token-budget branch splitter -- a differently-worded caption on a
    # re-run could shift branch boundaries and total node count.
    # Memoizing by the image's own byte hash makes two calls with the
    # SAME bytes return the SAME caption without a second vision call.
    import app.main as main_module

    class _FakeVisionClient:
        def __init__(self) -> None:
            self.calls = 0

        def describe_figure(self, image_bytes: bytes, mime_type: str | None = None, document_context: str | None = None, user_id: str | None = None) -> FigureVisionResult | None:
            self.calls += 1
            return FigureVisionResult(significant=True, description=f"caption #{self.calls}")

    class _FakeRedis:
        def __init__(self) -> None:
            self._store: dict[str, bytes] = {}

        def get(self, key):
            return self._store.get(key)

        def set(self, key, value, ex=None):
            self._store[key] = value.encode("utf-8") if isinstance(value, str) else value

        def publish(self, channel, payload):
            pass

    fake_client = _FakeVisionClient()
    fake_redis = _FakeRedis()
    monkeypatch.setattr(main_module, "_vision_client", fake_client)
    monkeypatch.setattr(main_module, "_narration_redis", fake_redis)

    first = main_module._describe_figure(b"identical-image-bytes")
    second = main_module._describe_figure(b"identical-image-bytes")

    assert first == "caption #1"
    assert second == "caption #1"  # cache hit -- NOT "caption #2"
    assert fake_client.calls == 1


def test_describe_figure_calls_the_vision_client_again_for_genuinely_different_images(monkeypatch) -> None:
    import app.main as main_module

    class _FakeVisionClient:
        def __init__(self) -> None:
            self.calls = 0

        def describe_figure(self, image_bytes: bytes, mime_type: str | None = None, document_context: str | None = None, user_id: str | None = None) -> FigureVisionResult | None:
            self.calls += 1
            return FigureVisionResult(significant=True, description=f"caption #{self.calls}")

    class _FakeRedis:
        def __init__(self) -> None:
            self._store: dict[str, bytes] = {}

        def get(self, key):
            return self._store.get(key)

        def set(self, key, value, ex=None):
            self._store[key] = value.encode("utf-8") if isinstance(value, str) else value

        def publish(self, channel, payload):
            pass

    fake_client = _FakeVisionClient()
    monkeypatch.setattr(main_module, "_vision_client", fake_client)
    monkeypatch.setattr(main_module, "_narration_redis", _FakeRedis())

    first = main_module._describe_figure(b"image-one")
    second = main_module._describe_figure(b"image-two")

    assert first == "caption #1"
    assert second == "caption #2"
    assert fake_client.calls == 2


def test_describe_figure_strips_a_degenerate_continuation_from_an_already_cached_entry(monkeypatch) -> None:
    # Real gap found during review, on top of the repetition-loop fix
    # itself: a garbled cache entry written BEFORE this fix landed (same
    # PROMPT_VERSION -- this fix changed parsing logic, not the prompt/
    # schema text PROMPT_VERSION is derived from, so the cache key alone
    # would NOT invalidate it) must still self-heal on the very next read,
    # not silently keep serving broken text for the rest of its 1-week TTL.
    # A vision client that fails the test if called proves this is
    # genuinely served from the cache HIT path, not a fresh call.
    import json

    import app.main as main_module

    class _FailIfCalledVisionClient:
        def describe_figure(self, image_bytes: bytes, mime_type: str | None = None, document_context: str | None = None, user_id: str | None = None):
            raise AssertionError("must be served from cache, not a fresh vision call")

    class _FakeRedis:
        def __init__(self, store: dict[str, bytes]) -> None:
            self._store = store

        def get(self, key):
            return self._store.get(key)

        def set(self, key, value, ex=None):
            raise AssertionError("a cache HIT must never re-write the entry")

        def publish(self, channel, payload):
            pass

    real_caption = "This drawing shows a rectangular panel measuring 46.43 by 40.72 units."
    garbled = real_caption + ".} } " + real_caption + " } } ```json {\"broken"
    cache_key = main_module._figure_caption_cache_key(b"already-cached-image", None)
    store = {cache_key: json.dumps({"significant": True, "description": garbled}).encode("utf-8")}

    monkeypatch.setattr(main_module, "_vision_client", _FailIfCalledVisionClient())
    monkeypatch.setattr(main_module, "_narration_redis", _FakeRedis(store))

    result = main_module._describe_figure(b"already-cached-image")

    assert result is not None
    assert "```" not in result
    # Everything up to (not including) the code fence survives -- this
    # backstop targets the unambiguous fence artifact specifically, not
    # full duplicate-content dedup (see _strip_degenerate_continuation's
    # own docstring for why: frequency_penalty is the actual root-cause
    # mitigation for milder in-content repetition).
    assert result == real_caption + ".} } " + real_caption + " } }"


def test_figure_caption_cache_key_includes_prompt_version() -> None:
    # Real gap found during review: a prompt wording change (which happened
    # twice in one real debugging session already) must invalidate old
    # cache entries -- otherwise a real deployment would keep serving
    # captions generated under the OLD prompt for up to the full 1-week
    # TTL, with the fix looking shipped but not observably active. Asserts
    # the key genuinely contains PROMPT_VERSION, not just that changing it
    # would change the key (which the property test in
    # test_orchestrator_vision_client.py already covers for PROMPT_VERSION
    # itself).
    key = _figure_caption_cache_key(b"some-image-bytes", "some document context")
    assert PROMPT_VERSION in key


def test_describe_figure_cache_key_does_not_collide_across_different_document_contexts(monkeypatch) -> None:
    # Regression guard for _figure_caption_cache_key's document_context
    # component: the SAME image bytes embedded in two DIFFERENT documents
    # can legitimately get two DIFFERENT significance verdicts once
    # document context is a real input to that judgment (an image-only
    # cache key would silently leak document A's verdict onto document B).
    # This fake returns a caption keyed off the document_context it was
    # actually handed, so a false cache hit would surface as the WRONG
    # caption coming back, not just a missing second call.
    import app.main as main_module

    class _FakeVisionClient:
        def __init__(self) -> None:
            self.calls = 0

        def describe_figure(self, image_bytes: bytes, mime_type: str | None = None, document_context: str | None = None, user_id: str | None = None) -> FigureVisionResult | None:
            self.calls += 1
            return FigureVisionResult(significant=True, description=f"caption for {document_context}")

    class _FakeRedis:
        def __init__(self) -> None:
            self._store: dict[str, bytes] = {}

        def get(self, key):
            return self._store.get(key)

        def set(self, key, value, ex=None):
            self._store[key] = value.encode("utf-8") if isinstance(value, str) else value

        def publish(self, channel, payload):
            pass

    fake_client = _FakeVisionClient()
    monkeypatch.setattr(main_module, "_vision_client", fake_client)
    monkeypatch.setattr(main_module, "_narration_redis", _FakeRedis())

    first = main_module._describe_figure(b"same-image-bytes", document_context="Document A: an avionics ICD.")
    second = main_module._describe_figure(b"same-image-bytes", document_context="Document B: a plumbing schematic.")

    # NOT served from a cross-document cache collision -- each document_context
    # gets its own live vision call and its own correct caption back.
    assert first == "caption for Document A: an avionics ICD."
    assert second == "caption for Document B: a plumbing schematic."
    assert fake_client.calls == 2


def test_describe_figure_degrades_to_a_live_call_when_the_cache_read_or_write_fails(monkeypatch) -> None:
    # Fail-open, same convention as the dedup check (dry-run.service.ts) and
    # KB retrieval (retrieval.py) -- a cache outage must never break figure
    # captioning, only make it non-cached for that call.
    import app.main as main_module

    class _FakeVisionClient:
        def describe_figure(self, image_bytes: bytes, mime_type: str | None = None, document_context: str | None = None, user_id: str | None = None) -> FigureVisionResult | None:
            return FigureVisionResult(significant=True, description="live caption")

    class _ExplodingRedis:
        def get(self, key):
            raise ConnectionError("redis is down")

        def set(self, key, value, ex=None):
            raise ConnectionError("redis is down")

    monkeypatch.setattr(main_module, "_vision_client", _FakeVisionClient())
    monkeypatch.setattr(main_module, "_narration_redis", _ExplodingRedis())

    assert main_module._describe_figure(b"some-image-bytes") == "live caption"


def test_describe_figure_never_caches_a_blank_result(monkeypatch) -> None:
    # Fail-open protection: an unparseable/failed vision response (None,
    # see FigureVisionResult's own doc comment) must NOT be cached -- it
    # would otherwise replay that same "unknown" answer for a full week
    # (_FIGURE_CAPTION_CACHE_TTL_SECONDS), masking whatever caused the
    # original failure. A genuinely parsed result on retry must succeed
    # normally and get cached like any other real verdict.
    import app.main as main_module

    class _FakeVisionClient:
        def __init__(self) -> None:
            self.calls = 0

        def describe_figure(self, image_bytes: bytes, mime_type: str | None = None, document_context: str | None = None, user_id: str | None = None) -> FigureVisionResult | None:
            self.calls += 1
            return None if self.calls == 1 else FigureVisionResult(significant=True, description="a real caption on retry")

    class _FakeRedis:
        def __init__(self) -> None:
            self._store: dict[str, bytes] = {}

        def get(self, key):
            return self._store.get(key)

        def set(self, key, value, ex=None):
            self._store[key] = value.encode("utf-8") if isinstance(value, str) else value

        def publish(self, channel, payload):
            pass

    fake_client = _FakeVisionClient()
    monkeypatch.setattr(main_module, "_vision_client", fake_client)
    monkeypatch.setattr(main_module, "_narration_redis", _FakeRedis())

    first = main_module._describe_figure(b"same-image-bytes")
    second = main_module._describe_figure(b"same-image-bytes")

    assert first is None
    # NOT served from a poisoned cache -- the None result was never
    # written, so the second call goes live again and self-heals.
    assert second == "a real caption on retry"
    assert fake_client.calls == 2


def test_describe_figure_caches_a_genuine_insignificant_verdict(monkeypatch) -> None:
    # The mirror image of the never-cache-a-blank test above, and the reason
    # the "ambiguous" cases (unparseable response, or significant=True with a
    # blank description) are collapsed into None rather than into a
    # significant=False result: a REAL significant=False verdict is a
    # confident, deliberate answer and MUST be cached like any other, so a
    # re-run of the same document doesn't re-pay a vision call per decorative
    # figure. Both calls use identical image bytes AND identical
    # document_context, so they share one cache key.
    import app.main as main_module
    from app.vision.orchestrator_vision_client import FigureVisionResult

    class _FakeVisionClient:
        def __init__(self) -> None:
            self.calls = 0

        def describe_figure(self, image_bytes, mime_type=None, document_context=None, user_id=None):
            self.calls += 1
            return FigureVisionResult(significant=False, description="")

    fake_client = _FakeVisionClient()
    monkeypatch.setattr(main_module, "_vision_client", fake_client)
    monkeypatch.setattr(main_module, "_narration_redis", _FakeRedis())

    first = main_module._describe_figure(b"decorative-image-bytes", document_context="An avionics ICD.")
    second = main_module._describe_figure(b"decorative-image-bytes", document_context="An avionics ICD.")

    assert first is None
    assert second is None
    assert fake_client.calls == 1  # second call served from cache, no live re-query


def test_describe_figure_returns_none_when_vision_model_judges_figure_insignificant(monkeypatch) -> None:
    import app.main as main_module
    from app.vision.orchestrator_vision_client import FigureVisionResult

    class _FakeVisionClient:
        def describe_figure(self, image_bytes, mime_type=None, document_context=None, user_id=None):
            return FigureVisionResult(significant=False, description="A decorative rectangular block.")

    monkeypatch.setattr(main_module, "_vision_client", _FakeVisionClient())
    monkeypatch.setattr(main_module, "_narration_redis", _FakeRedis())

    result = main_module._describe_figure(b"fake-image-bytes")
    assert result is None


def test_describe_figure_returns_the_description_when_vision_model_judges_figure_significant(monkeypatch) -> None:
    import app.main as main_module
    from app.vision.orchestrator_vision_client import FigureVisionResult

    class _FakeVisionClient:
        def describe_figure(self, image_bytes, mime_type=None, document_context=None, user_id=None):
            return FigureVisionResult(significant=True, description="A power interface wiring diagram.")

    monkeypatch.setattr(main_module, "_vision_client", _FakeVisionClient())
    monkeypatch.setattr(main_module, "_narration_redis", _FakeRedis())

    result = main_module._describe_figure(b"fake-image-bytes")
    assert result == "A power interface wiring diagram."


def test_describe_figure_forwards_document_context_to_the_vision_client(monkeypatch) -> None:
    import app.main as main_module
    from app.vision.orchestrator_vision_client import FigureVisionResult

    captured = {}

    class _FakeVisionClient:
        def describe_figure(self, image_bytes, mime_type=None, document_context=None, user_id=None):
            captured["document_context"] = document_context
            return FigureVisionResult(significant=True, description="ok")

    monkeypatch.setattr(main_module, "_vision_client", _FakeVisionClient())
    monkeypatch.setattr(main_module, "_narration_redis", _FakeRedis())

    main_module._describe_figure(b"fake-image-bytes", document_context="An avionics ICD.")
    assert captured["document_context"] == "An avionics ICD."


def test_ingest_takes_identity_and_the_document_overview_from_headers_only(monkeypatch) -> None:
    # The overview and the user used to be query keys; they travel in headers so they never reach an access log. A mismatched name would
    # silently swallow the value, so the signature is pinned: no identity query parameters, and no `document_context` either.
    import inspect

    import app.main as main_module

    sig = inspect.signature(main_module.ingest)
    assert "primer_context" not in sig.parameters and "user_id" not in sig.parameters
    assert "document_context" not in sig.parameters
    assert {"x_user_id", "x_primer_context"} <= set(sig.parameters)


def test_ingest_publishes_describe_figure_and_parse_done_narration_when_job_id_is_given(monkeypatch) -> None:
    # job_id is optional (absent for /jobs/preview's own call, which has no
    # job yet to narrate for) -- when a real job submission supplies it,
    # /ingest must publish real parse-stage progress on the same channel
    # /level's SURVEY narration already uses, instead of leaving the parse
    # stage a total black box for however long figure captioning takes.
    import app.main as main_module

    class _FakeVisionClient:
        def describe_figure(self, image_bytes: bytes, mime_type: str | None = None, document_context: str | None = None, user_id: str | None = None) -> FigureVisionResult | None:
            return FigureVisionResult(significant=True, description="a real caption")

    class _FakeRedis:
        def __init__(self) -> None:
            self.published: list[str] = []

        def set(self, key, value, ex=None):
            pass

        def rpush(self, key, value):
            pass

        def ltrim(self, key, start, end):
            pass

        def expire(self, key, seconds):
            pass

        def publish(self, channel, payload):
            self.published.append(payload)

    fake_redis = _FakeRedis()
    monkeypatch.setattr(main_module, "_vision_client", _FakeVisionClient())
    monkeypatch.setattr(main_module, "_narration_redis", fake_redis)

    _upload("sample.docx", FIXTURES / "sample.docx", mode="survey", job_id="job-narration-test")

    import json

    payloads = [json.loads(p) for p in fake_redis.published]
    phases = [p["phase"] for p in payloads]
    assert "describe_figure" in phases
    assert "parse_done" in phases

    # Two describe_figure events for this ONE figure -- an in-flight signal
    # (fired before the vision call, figureCaption null) and a real-result
    # signal (fired after, carrying the actual caption text). Without the
    # second, a figure-heavy document showed the same generic "figure N so
    # far" line for every figure with zero differentiating content
    # (confirmed live: 7 identical rows for 7 real figures).
    figure_events = [p for p in payloads if p["phase"] == "describe_figure"]
    assert len(figure_events) == 2
    assert figure_events[0]["figureIndex"] == 1
    assert figure_events[0]["figureCaption"] is None
    assert figure_events[1]["figureIndex"] == 1
    assert figure_events[1]["figureCaption"] == "a real caption"
    assert "a real caption" in figure_events[1]["message"]

    # Real gap this guards against (item 16): a .docx upload should now run
    # a cheap structural pre-count ahead of the real walk, so the in-flight
    # describe_figure event's message carries a real "figure 1 of 1"
    # denominator instead of the old open-ended "figure 1 so far" -- sample.docx
    # has exactly one figure, so index and total are both 1.
    assert figure_events[0]["figureCount"] == 1
    assert "figure 1 of 1" in figure_events[0]["message"]

    parse_done_event = next(p for p in payloads if p["phase"] == "parse_done")
    assert parse_done_event["tableCount"] == 1  # sample.docx has 1 table
    assert parse_done_event["figureCount"] == 1

    # Real bug this guards against (confirmed live): tableRowCount was
    # previously derived from rows_from_table (the ID-hierarchy-QUALIFYING
    # rows only, legitimately [] for a table whose content isn't ID-shaped)
    # instead of the table's own real raw row count -- every non-qualifying
    # table read as "0 row(s)", indistinguishable from a genuinely empty or
    # mis-parsed one. Ground truth comes from parsing the SAME fixture
    # directly, not a hardcoded number, so this stays correct if the fixture
    # ever changes.
    from app.dispatch import dispatch_by_extension

    real_doc = dispatch_by_extension("sample.docx", str(FIXTURES / "sample.docx"))
    table_events = [p for p in payloads if p["phase"] == "table_detected"]
    assert len(table_events) == 1
    assert table_events[0]["tableRowCount"] == len(real_doc.tables[0].cells)
    assert table_events[0]["tableRowCount"] > 0
    # table_detected doesn't need a pre-count -- it fires from a post-parse
    # loop over the full table list, so position/total are always known.
    assert table_events[0]["tableIndex"] == 1
    assert table_events[0]["tableCount"] == 1
    assert "table 1 of 1" in table_events[0]["message"]


def test_ingest_pptx_publishes_a_real_figure_count_denominator(monkeypatch) -> None:
    # SWEEP.md item 3: the docx-only precount above previously left pptx on
    # the old open-ended "figure N so far" narration -- extends the same
    # exact structural pre-pass to pptx.
    import app.main as main_module

    class _FakeVisionClient:
        def describe_figure(self, image_bytes: bytes, mime_type: str | None = None, document_context: str | None = None, user_id: str | None = None) -> FigureVisionResult | None:
            return FigureVisionResult(significant=True, description="a red square")

    class _FakeRedis:
        def __init__(self) -> None:
            self.published: list[str] = []

        def set(self, key, value, ex=None):
            pass

        def rpush(self, key, value):
            pass

        def ltrim(self, key, start, end):
            pass

        def expire(self, key, seconds):
            pass

        def publish(self, channel, payload):
            self.published.append(payload)

    monkeypatch.setattr(main_module, "_vision_client", _FakeVisionClient())
    fake_redis = _FakeRedis()
    monkeypatch.setattr(main_module, "_narration_redis", fake_redis)

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    image_buffer = io.BytesIO()
    Image.new("RGB", (20, 20), color="red").save(image_buffer, format="PNG")
    image_buffer.seek(0)
    slide.shapes.add_picture(image_buffer, Inches(1), Inches(1))
    path = Path(tempfile.mkdtemp()) / "one_figure.pptx"
    prs.save(str(path))

    _upload("one_figure.pptx", path, mode="survey", job_id="job-pptx-precount")

    import json

    payloads = [json.loads(p) for p in fake_redis.published]
    figure_events = [p for p in payloads if p["phase"] == "describe_figure"]
    assert figure_events[0]["figureCount"] == 1
    assert "figure 1 of 1" in figure_events[0]["message"]


def test_ingest_image_publishes_a_figure_count_of_one_with_no_precount_walk(monkeypatch) -> None:
    # A standalone image upload is unconditionally exactly one Figure (see
    # image_loader.load_image's own docstring) -- no structural walk needed.
    import app.main as main_module

    class _FakeVisionClient:
        def describe_figure(self, image_bytes: bytes, mime_type: str | None = None, document_context: str | None = None, user_id: str | None = None) -> FigureVisionResult | None:
            return FigureVisionResult(significant=True, description="a real caption")

    class _FakeRedis:
        def __init__(self) -> None:
            self.published: list[str] = []

        def set(self, key, value, ex=None):
            pass

        def rpush(self, key, value):
            pass

        def ltrim(self, key, start, end):
            pass

        def expire(self, key, seconds):
            pass

        def publish(self, channel, payload):
            self.published.append(payload)

    monkeypatch.setattr(main_module, "_vision_client", _FakeVisionClient())
    fake_redis = _FakeRedis()
    monkeypatch.setattr(main_module, "_narration_redis", fake_redis)

    _upload("sample.png", FIXTURES / "sample.png", mode="survey", job_id="job-image-precount")

    import json

    payloads = [json.loads(p) for p in fake_redis.published]
    figure_events = [p for p in payloads if p["phase"] == "describe_figure"]
    assert figure_events[0]["figureCount"] == 1
    assert "figure 1 of 1" in figure_events[0]["message"]


def test_ingest_pdf_publishes_the_exact_figure_count_denominator(monkeypatch) -> None:
    # pdf_with_figure.pdf has exactly one embedded image XObject. This used
    # to be an APPROXIMATE precount via count_image_xobjects (a cheap pass
    # avoiding a second, expensive load_pdf walk) -- since figure captioning
    # moved to a post-parse pass (SWEEP 8b), the denominator is now simply
    # the number of figures the single real parse actually collected, exact
    # for every format including pdf.
    import app.main as main_module

    class _FakeVisionClient:
        def describe_figure(self, image_bytes: bytes, mime_type: str | None = None, document_context: str | None = None, user_id: str | None = None) -> FigureVisionResult | None:
            return FigureVisionResult(significant=True, description="a real caption")

    class _FakeRedis:
        def __init__(self) -> None:
            self.published: list[str] = []

        def set(self, key, value, ex=None):
            pass

        def rpush(self, key, value):
            pass

        def ltrim(self, key, start, end):
            pass

        def expire(self, key, seconds):
            pass

        def publish(self, channel, payload):
            self.published.append(payload)

    monkeypatch.setattr(main_module, "_vision_client", _FakeVisionClient())
    fake_redis = _FakeRedis()
    monkeypatch.setattr(main_module, "_narration_redis", fake_redis)

    _upload("pdf_with_figure.pdf", FIXTURES / "pdf_with_figure.pdf", mode="survey", job_id="job-pdf-precount")

    import json

    payloads = [json.loads(p) for p in fake_redis.published]
    figure_events = [p for p in payloads if p["phase"] == "describe_figure"]
    assert figure_events[0]["figureCount"] == 1
    assert "figure 1 of 1" in figure_events[0]["message"]


def test_ingest_publishes_no_narration_when_job_id_is_omitted(monkeypatch) -> None:
    # /jobs/preview's own call never supplies job_id -- must not publish
    # anything (there is no job to attach the narration to).
    import app.main as main_module

    class _FakeVisionClient:
        def describe_figure(self, image_bytes: bytes, mime_type: str | None = None, document_context: str | None = None, user_id: str | None = None) -> FigureVisionResult | None:
            return FigureVisionResult(significant=True, description="a real caption")

    class _FakeRedis:
        def __init__(self) -> None:
            self.published: list[str] = []

        def set(self, key, value, ex=None):
            pass

        def rpush(self, key, value):
            pass

        def ltrim(self, key, start, end):
            pass

        def expire(self, key, seconds):
            pass

        def publish(self, channel, payload):
            self.published.append(payload)

    fake_redis = _FakeRedis()
    monkeypatch.setattr(main_module, "_vision_client", _FakeVisionClient())
    monkeypatch.setattr(main_module, "_narration_redis", fake_redis)

    _upload("sample.docx", FIXTURES / "sample.docx", mode="survey")

    assert fake_redis.published == []


def test_ingest_baseline_mode_never_captions_even_when_vision_is_configured(monkeypatch) -> None:
    # BASELINE is contractually verbatim/no-AI: a figure caption is
    # model-authored text, and must never reach a contract-baseline
    # System.Description regardless of skip_vision or whether a vision
    # client is available. This is the mode gate
    # itself -- distinct from skip_vision, which is a latency optimization
    # for /jobs/preview only. Default mode ("baseline", per /ingest's own
    # signature) must behave identically to an explicit mode="baseline".
    import app.main as main_module

    class _FakeVisionClient:
        def __init__(self) -> None:
            self.calls = 0

        def describe_figure(self, image_bytes: bytes, mime_type: str | None = None, user_id: str | None = None) -> FigureVisionResult | None:
            self.calls += 1
            return FigureVisionResult(significant=True, description="a description that must never appear in baseline")

    fake_client = _FakeVisionClient()
    monkeypatch.setattr(main_module, "_vision_client", fake_client)

    body = _upload("sample.docx", FIXTURES / "sample.docx")  # mode omitted -> default baseline
    figure_blocks = [b for b in body["blocks"] if b["kind"] == "figure"]
    assert figure_blocks[0]["text"] == "[FIGURE: none]"
    assert fake_client.calls == 0

    fake_client.calls = 0
    body = _upload("sample.docx", FIXTURES / "sample.docx", mode="baseline")  # explicit, same result
    figure_blocks = [b for b in body["blocks"] if b["kind"] == "figure"]
    assert figure_blocks[0]["text"] == "[FIGURE: none]"
    assert fake_client.calls == 0


def test_ingest_unrecognized_mode_value_fails_safe_to_baseline(monkeypatch) -> None:
    # Fail-safe direction matches /level's own `if mode == "survey"` fork
    # (main.py) -- anything other than the exact string "survey" must never
    # open the vision gate.
    import app.main as main_module

    class _FakeVisionClient:
        def __init__(self) -> None:
            self.calls = 0

        def describe_figure(self, image_bytes: bytes, mime_type: str | None = None, user_id: str | None = None) -> FigureVisionResult | None:
            self.calls += 1
            return FigureVisionResult(significant=True, description="should never be reached")

    fake_client = _FakeVisionClient()
    monkeypatch.setattr(main_module, "_vision_client", fake_client)

    body = _upload("sample.docx", FIXTURES / "sample.docx", mode="Survey")  # wrong case, not "survey"
    figure_blocks = [b for b in body["blocks"] if b["kind"] == "figure"]
    assert figure_blocks[0]["text"] == "[FIGURE: none]"
    assert fake_client.calls == 0


def test_ingest_then_level_baseline_never_touches_a_configured_vision_client() -> None:
    # Executable version of the "BASELINE is no-AI" invariant that
    # config.yaml's enabled_in_baseline: false documents but nothing reads
    # (see docs/architecture audit): drives the real pipeline shape
    # (/ingest?mode=baseline -> /level?mode=baseline) end to end over a
    # figure-bearing document with a real vision client monkeypatched to
    # raise on any call. If BASELINE ever reopens the vision gate, this
    # fails loudly instead of silently emitting a model-authored caption.
    import app.main as main_module

    class _ExplodingVisionClient:
        def describe_figure(self, image_bytes: bytes, mime_type: str | None = None, user_id: str | None = None) -> FigureVisionResult | None:
            raise AssertionError("BASELINE must never call describe_figure")

    monkeypatch_target = main_module._vision_client
    main_module._vision_client = _ExplodingVisionClient()
    try:
        with open(FIXTURES / "sample.docx", "rb") as f:
            ingest_response = client.post(
                "/ingest",
                params={"mode": "baseline"},
                files={"file": ("sample.docx", f, "application/octet-stream")},
            )
        assert ingest_response.status_code == 200
        doc = ingest_response.json()

        level_response = client.post("/level", params={"mode": "baseline"}, json=doc)
        assert level_response.status_code == 200
        figure_blocks = [n for n in level_response.json()["nodes"] if "[FIGURE:" in n["description"]]
        assert all("none" in n["description"] for n in figure_blocks)
    finally:
        main_module._vision_client = monkeypatch_target


def test_level_survey_reports_the_real_error_detail_not_a_bare_500() -> None:
    # Confirmed real UX gap (user-reported): a failed SURVEY job's audit
    # trail only ever showed "Ingestion /level request failed: 500" --
    # accurate but useless, since Starlette's default (no handler
    # registered) discards the real exception message into a generic
    # '{"detail": "Internal Server Error"}' body. The real message (e.g. a
    # malformed-JSON LLM response) must reach the response body so the
    # orchestrator's audit entry can show it.
    import app.main as main_module

    def _raise_value_error(*args: object, **kwargs: object) -> None:
        raise ValueError("malformed JSON in LLM response content: 'Based'")

    original = main_module._build_survey_tree_with_scoring
    main_module._build_survey_tree_with_scoring = _raise_value_error
    try:
        response = client.post(
            "/level",
            params={"mode": "survey"},
            json={"document": {"path": "x.txt", "format": "txt", "sha256": "h"}, "blocks": [], "tables": [], "figures": []},
        )
        # ValueError means "an upstream dependency (the LLM) misbehaved" --
        # 502 Bad Gateway, not a generic 500, and the real message survives.
        assert response.status_code == 502
        assert response.json()["detail"] == "malformed JSON in LLM response content: 'Based'"
    finally:
        main_module._build_survey_tree_with_scoring = original


def test_level_survey_reports_a_genuine_internal_error_as_500() -> None:
    import app.main as main_module

    def _raise_runtime_error(*args: object, **kwargs: object) -> None:
        raise RuntimeError("unexpected internal failure")

    original = main_module._build_survey_tree_with_scoring
    main_module._build_survey_tree_with_scoring = _raise_runtime_error
    # TestClient's default raise_server_exceptions=True re-raises ANY
    # unhandled-exception-shaped ASGI signal for debugging visibility, even
    # when a registered exception_handler already produced a real response
    # -- a real production ASGI server (uvicorn) has no such re-raise
    # behavior and would return the handler's response exactly like the
    # ValueError case above did. A dedicated client with this off is the
    # standard way to test a bare Exception handler through the full stack.
    lenient_client = TestClient(main_module.app, raise_server_exceptions=False)
    try:
        response = lenient_client.post(
            "/level",
            params={"mode": "survey"},
            json={"document": {"path": "x.txt", "format": "txt", "sha256": "h"}, "blocks": [], "tables": [], "figures": []},
        )
        assert response.status_code == 500
        assert response.json()["detail"] == "unexpected internal failure"
    finally:
        main_module._build_survey_tree_with_scoring = original


def test_build_survey_tree_with_scoring_passes_the_client_as_title_tag_client() -> None:
    # Confirmed real gap: main.py's /level call site never actually passed
    # a title_tag_client at all, so every deterministically-built node kept
    # its raw extract_baseline_title-derived title and empty tags in
    # production, even though the title/tag enhancement pass was fully
    # built and tested. OrchestratorLlmClient implements the TitleTagClient
    # protocol directly, so it's the SAME client already passed as
    # analysis_client/scoring_client/etc above.
    import app.main as main_module
    from app.canonical.model import CanonicalDoc, Document, DocumentFormat
    from app.leveling.tree_builder import WorkItemTree

    captured: dict[str, object] = {}

    def _capture(doc, **kwargs):
        captured.update(kwargs)
        return WorkItemTree(nodes=[], warnings=[], kb_citations=[], trace_links=[])

    original = main_module.build_survey_tree
    main_module.build_survey_tree = _capture
    try:
        fake_client = object()
        main_module._build_survey_tree_with_scoring(
            doc=CanonicalDoc(
                document=Document(path="x.txt", format=DocumentFormat.TXT, sha256="h"),
                blocks=[], tables=[], figures=[],
            ),
            client=fake_client,
            user_id=None,
            link_type_catalog=None,
            job_id=None,
        )
        assert captured["title_tag_client"] is fake_client
    finally:
        main_module.build_survey_tree = original


def test_ingest_unsupported_extension_returns_400() -> None:
    import tempfile

    with tempfile.NamedTemporaryFile(suffix=".exe", mode="wb", delete=False) as f:
        f.write(b"not a real document")
        temp_path = Path(f.name)

    with open(temp_path, "rb") as f:
        response = client.post(
            "/ingest", files={"file": (temp_path.name, f, "application/octet-stream")}
        )
    assert response.status_code == 400
    temp_path.unlink()


def test_sniff_image_mime_type_detects_jpeg() -> None:
    assert _sniff_image_mime_type(b"\xff\xd8\xff\xe0rest of jpeg") == "image/jpeg"


def test_sniff_image_mime_type_detects_png() -> None:
    assert _sniff_image_mime_type(b"\x89PNG\r\n\x1a\nrest of png") == "image/png"


def test_sniff_image_mime_type_detects_gif() -> None:
    assert _sniff_image_mime_type(b"GIF89arest of gif") == "image/gif"


def test_sniff_image_mime_type_detects_bmp() -> None:
    assert _sniff_image_mime_type(b"BMrest of bmp") == "image/bmp"


def test_sniff_image_mime_type_detects_webp() -> None:
    assert _sniff_image_mime_type(b"RIFF\x00\x00\x00\x00WEBPrest") == "image/webp"


def test_sniff_image_mime_type_detects_tiff() -> None:
    assert _sniff_image_mime_type(b"II*\x00rest of tiff") == "image/tiff"


def test_sniff_image_mime_type_returns_none_for_unrecognized_bytes() -> None:
    assert _sniff_image_mime_type(b"not a real image format") is None


def test_rasterize_for_vision_passes_through_a_real_raster_format_unchanged() -> None:
    png_bytes = b"\x89PNG\r\n\x1a\n" + b"rest of png"
    rasterized, mime_type = _rasterize_for_vision(png_bytes)
    assert rasterized == png_bytes
    assert mime_type == "image/png"


def test_rasterize_for_vision_converts_a_vector_format_via_emf_converter(monkeypatch) -> None:
    # Real gap this guards against (confirmed live): compress_figure_for_
    # storage (Pillow-only) silently failed for an EMF/WMF figure that
    # _describe_figure captioned successfully via this exact conversion --
    # a figure could be fully described yet never become storable/viewable.
    # This is the shared helper both call sites must now go through.
    import app.main as main_module

    fake_emf_bytes = b"fake emf bytes"
    fake_png_bytes = b"\x89PNG\r\n\x1a\nconverted"
    monkeypatch.setattr(main_module, "sniff_vector_image_extension", lambda _b: ".emf")
    monkeypatch.setattr(main_module, "convert_vector_image_to_png", lambda _b, _ext: fake_png_bytes)

    rasterized, mime_type = _rasterize_for_vision(fake_emf_bytes)
    assert rasterized == fake_png_bytes
    assert mime_type == "image/png"


def test_rasterize_for_vision_falls_back_to_original_bytes_when_conversion_fails(monkeypatch) -> None:
    import app.main as main_module

    fake_emf_bytes = b"fake emf bytes, conversion will fail"
    monkeypatch.setattr(main_module, "sniff_vector_image_extension", lambda _b: ".emf")
    monkeypatch.setattr(main_module, "convert_vector_image_to_png", lambda _b, _ext: None)

    rasterized, mime_type = _rasterize_for_vision(fake_emf_bytes)
    assert rasterized == fake_emf_bytes  # unchanged -- same degrade-gracefully contract as before
    assert mime_type is None


def test_rasterize_for_vision_returns_original_bytes_for_a_genuinely_unrecognized_format() -> None:
    rasterized, mime_type = _rasterize_for_vision(b"not a real image at all")
    assert rasterized == b"not a real image at all"
    assert mime_type is None


def test_figure_storage_uses_the_same_rasterized_bytes_as_captioning_not_the_original(monkeypatch) -> None:
    # Direct regression test for the confirmed real bug (reported live,
    # "not the first time it occurs"): sample.docx's one embedded figure is
    # a real raster PNG, so it wouldn't naturally exercise the vector-
    # conversion path -- sniff_vector_image_extension is force-monkeypatched
    # to claim EVERY figure is an EMF regardless of its real bytes, making
    # this test deterministic without needing a real EMF-embedded fixture
    # or a real soffice conversion.
    import app.main as main_module

    fake_converted_bytes = b"\x89PNG\r\n\x1a\nconverted-for-storage"
    captured_store_calls: list[tuple[str, str, bytes, str]] = []

    class _FakeVisionClient:
        def describe_figure(self, image_bytes: bytes, mime_type: str | None = None, document_context: str | None = None, user_id: str | None = None) -> FigureVisionResult | None:
            return FigureVisionResult(significant=True, description="a real caption")

    class _FakeFigureFileClient:
        def store(self, job_id: str, figure_id: str, image_bytes: bytes, mime_type: str) -> None:
            captured_store_calls.append((job_id, figure_id, image_bytes, mime_type))

    def _fake_compress_for_storage(image_bytes: bytes):
        # Returns whatever it was handed, unmodified -- lets this test
        # assert exactly which bytes reached the storage call, without also
        # depending on real PNG re-encoding/Pillow behavior.
        return (image_bytes, "image/png")

    monkeypatch.setattr(main_module, "_vision_client", _FakeVisionClient())
    monkeypatch.setattr(main_module, "_figure_file_client", _FakeFigureFileClient())
    # sample.docx's one embedded figure is a real PNG, so _sniff_image_mime_
    # type would correctly identify it and short-circuit BEFORE ever
    # reaching the vector-extension check below -- forced to None here so
    # this test exercises the vector-conversion branch deterministically,
    # regardless of the fixture's real (raster) content.
    monkeypatch.setattr(main_module, "_sniff_image_mime_type", lambda _b: None)
    monkeypatch.setattr(main_module, "sniff_vector_image_extension", lambda _b: ".emf")
    convert_calls: list[bytes] = []

    def _fake_convert(image_bytes: bytes, _ext: str) -> bytes:
        convert_calls.append(image_bytes)
        return fake_converted_bytes

    monkeypatch.setattr(main_module, "convert_vector_image_to_png", _fake_convert)
    monkeypatch.setattr(main_module, "compress_figure_for_storage", _fake_compress_for_storage)

    _upload("sample.docx", FIXTURES / "sample.docx", mode="survey", job_id="job-figure-storage-test")

    assert len(captured_store_calls) == 1
    _job_id, _figure_id, stored_bytes, stored_mime_type = captured_store_calls[0]
    assert stored_bytes == fake_converted_bytes  # the RASTERIZED bytes, not sample.docx's original figure bytes
    assert stored_mime_type == "image/png"
    # Real gap this guards against: the storage call previously rasterized
    # its OWN copy of the same figure (a second, redundant conversion)
    # instead of reusing the bytes _describe_figure already computed for
    # captioning -- this figure must only ever be converted once.
    assert len(convert_calls) == 1


def test_figure_bytes_are_rasterized_exactly_once_per_figure(monkeypatch) -> None:
    # Guards the prior sweep fix ("Figure bytes rasterized twice per figure")
    # with a CALL COUNT, not just a bytes-identity assertion: the caption
    # call and the storage call must share one rasterization via
    # caption_status_out["rasterized_bytes"]. This is the assertion that
    # would catch a restructuring which split captioning and storage apart,
    # since the storage side would then have no side-channel to read and
    # would silently rasterize its own second copy.
    import app.main as main_module

    class _FakeVisionClient:
        def describe_figure(self, image_bytes: bytes, mime_type: str | None = None, document_context: str | None = None, user_id: str | None = None) -> FigureVisionResult | None:
            return FigureVisionResult(significant=True, description="a real caption")

    class _FakeFigureFileClient:
        def store(self, job_id: str, figure_id: str, image_bytes: bytes, mime_type: str) -> None:
            return None

    rasterize_calls: list[bytes] = []
    original_rasterize = main_module._rasterize_for_vision

    def counting_rasterize(image_bytes: bytes):
        rasterize_calls.append(image_bytes)
        return original_rasterize(image_bytes)

    monkeypatch.setattr(main_module, "_vision_client", _FakeVisionClient())
    monkeypatch.setattr(main_module, "_figure_file_client", _FakeFigureFileClient())
    monkeypatch.setattr(main_module, "_rasterize_for_vision", counting_rasterize)
    monkeypatch.setattr(main_module, "compress_figure_for_storage", lambda b: (b, "image/png"))

    result = _upload(
        "sample.docx", FIXTURES / "sample.docx", mode="survey",
        job_id="job-rasterize-once", primer_context="rasterize-once-call-count",
    )

    # sample.docx has exactly one figure; one rasterization total, shared by
    # the caption call and the storage call.
    assert len(result["figures"]) == 1
    assert len(rasterize_calls) == 1


def test_caption_and_store_one_figure_returns_a_complete_outcome(monkeypatch) -> None:
    # Direct unit test of the extracted work unit, independent of the route.
    import app.main as main_module

    class _FakeVisionClient:
        def describe_figure(self, image_bytes: bytes, mime_type: str | None = None, document_context: str | None = None, user_id: str | None = None) -> FigureVisionResult | None:
            return FigureVisionResult(significant=True, description="prose about the figure", title="Short Title")

    stored: list[tuple[str, str]] = []

    class _FakeFigureFileClient:
        def store(self, job_id: str, figure_id: str, image_bytes: bytes, mime_type: str) -> None:
            stored.append((job_id, figure_id))

    monkeypatch.setattr(main_module, "_vision_client", _FakeVisionClient())
    monkeypatch.setattr(main_module, "_figure_file_client", _FakeFigureFileClient())
    monkeypatch.setattr(main_module, "compress_figure_for_storage", lambda b: (b, "image/png"))

    outcome = main_module._caption_and_store_one_figure(
        "f0", b"\x89PNG\r\n\x1a\nfake", "unit-test-caption-and-store", "job-unit",
    )

    assert outcome.figure_id == "f0"
    assert outcome.caption == "prose about the figure"
    assert outcome.status is None
    assert outcome.title == "Short Title"
    assert outcome.storage_failed is False
    assert stored == [("job-unit", "f0")]


def test_caption_and_store_one_figure_flags_a_storage_failure_without_losing_the_caption(monkeypatch) -> None:
    import app.main as main_module

    class _FakeVisionClient:
        def describe_figure(self, image_bytes: bytes, mime_type: str | None = None, document_context: str | None = None, user_id: str | None = None) -> FigureVisionResult | None:
            return FigureVisionResult(significant=True, description="prose about the figure")

    class _RaisingFigureFileClient:
        def store(self, job_id: str, figure_id: str, image_bytes: bytes, mime_type: str) -> None:
            raise RuntimeError("storage is down")

    monkeypatch.setattr(main_module, "_vision_client", _FakeVisionClient())
    monkeypatch.setattr(main_module, "_figure_file_client", _RaisingFigureFileClient())
    monkeypatch.setattr(main_module, "compress_figure_for_storage", lambda b: (b, "image/png"))

    outcome = main_module._caption_and_store_one_figure(
        "f3", b"\x89PNG\r\n\x1a\nfake", "unit-test-storage-failure", "job-unit",
    )

    # Storage is best-effort: the caption still comes back, the failure is
    # reported as a flag for the caller to fold into
    # CanonicalDoc.failed_figure_storage_ids.
    assert outcome.caption == "prose about the figure"
    assert outcome.storage_failed is True
