"""Run once to (re)generate mock_prd_flat.docx: `python build_mock_prd_flat_docx.py`.

Mock stand-in for the real customer document Mock_PRD_Sensor (Not
Properly Aligned).docx (not committed here). Same logical content as
mock_prd_sensor.docx, but Heading 1 is NEVER used -- every section sits
at the SAME (Heading 2) level, flat, with no grouping tier above them at
all. This is the real document shape that exposed a confirmed live
defect in an earlier version of app.leveling.heading_outline: the Word
"Title" paragraph (indistinguishable from a real heading at the Block
level) was the document's own sole shallower heading, so it got named as
its own top-level Epic, wrapping every real Feature in the document under
one Epic literally titled after the document itself. See that module's
own docstring (S1/S2) for the fix.
"""
from pathlib import Path

from docx import Document as DocxDocument

doc = DocxDocument()

doc.add_heading("Product Requirements Document", level=0)  # python-docx's "Title" style
doc.add_paragraph("Mock / Anonymized Test Data")
doc.add_paragraph("Customer ID: 000")


def _add_requirements_table(doc: DocxDocument, prefix: str, rows: list[tuple[str, str]]) -> None:
    table = doc.add_table(rows=1, cols=3)
    header = table.rows[0].cells
    header[0].text, header[1].text, header[2].text = "#", "Parameter", "Requirement / Value"
    for i, (parameter, requirement) in enumerate(rows):
        cells = table.add_row().cells
        cells[0].text = f"{prefix}-{i:03d}#"
        cells[1].text = parameter
        cells[2].text = requirement


doc.add_heading("Synchronization", level=2)
_add_requirements_table(doc, "P-SENSOR", [
    ("Time sync source", "The system shall synchronize to Standard-XX within N ms."),
    ("Sync accuracy", "Timing accuracy shall be better than X microseconds."),
    ("Sync interval", "Re-synchronization shall occur every N seconds."),
])

doc.add_heading("Target Data report", level=2)
_add_requirements_table(doc, "P-TGT", [
    ("Report rate", "The system shall report target data at a rate of N Hz."),
    ("Report latency", "Report latency shall not exceed N ms end to end."),
])

doc.add_heading("Target tracking", level=2)
_add_requirements_table(doc, "P-TRK", [
    ("Target tracking count", "The system shall track up to N simultaneous targets."),
    ("Track update rate", "Track updates shall occur at a rate of N Hz."),
])

doc.add_heading("SW Safety", level=2)
_add_requirements_table(doc, "P-SAF", [
    ("Safety standard compliance", "The software shall comply with safety standard XYZ."),
    ("Watchdog reset time", "The watchdog shall reset the system within N ms of a fault."),
])

doc.add_heading("Connectors", level=2)
_add_requirements_table(doc, "P-CON", [
    ("Power connector type", "The power connector shall be of type XYZ."),
    ("Main connector pin count", "The main connector shall have N pins."),
])

out_path = Path(__file__).parent / "mock_prd_flat.docx"
doc.save(str(out_path))
print(f"wrote {out_path}")
