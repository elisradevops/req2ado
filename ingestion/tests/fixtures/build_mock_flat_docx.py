"""Run once to (re)generate mock_flat.docx: `python build_mock_flat_docx.py`.

Mock stand-in for a customer document with NO internal heading structure at
all — no Word Heading styles, no numbering prefixes anywhere. Drives the
"group into Epics/Features only when the document has no heading structure"
rule from the original request: today (pre-fix) a document like this
resolves every block's depth to None and produces ZERO work-item nodes.
After the fix, it must synthesize a single Epic -> Feature root pair with
every paragraph folded in as a Requirement child.
"""
from pathlib import Path

from docx import Document as DocxDocument

doc = DocxDocument()

doc.add_paragraph(
    "The system shall support at least four simultaneous audio channels "
    "without audible dropout."
)
doc.add_paragraph(
    "Startup power draw should stay under two watts to keep the thermal "
    "budget manageable in the enclosure."
)
doc.add_paragraph(
    "The unit shall survive a one-meter drop onto a hard surface without "
    "loss of function."
)

out_path = Path(__file__).parent / "mock_flat.docx"
doc.save(str(out_path))
print(f"wrote {out_path}")
