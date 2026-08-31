"""Run once to (re)generate multi_figure.docx: `python build_multi_figure_docx_fixture.py`.

Two-phase-figure-captioning plan (SWEEP 8b): needed a static fixture with at
least two DISTINCT embedded images (distinct bytes -- the loaders dedupe by
sha256, so two copies of the same image collapse to one figure). No existing
fixture had more than one embedded image, so this one was added, following
build_docx_fixture.py's own generation pattern exactly.
"""
import tempfile
from pathlib import Path

from docx import Document as DocxDocument
from PIL import Image

doc = DocxDocument()
doc.add_heading("Display", level=1)
doc.add_paragraph("The display shall support at least three picture presets.")
doc.add_paragraph("See the first diagram below for panel layout.")

with tempfile.TemporaryDirectory() as tmp_dir:
    image_one_path = Path(tmp_dir) / "synthetic_red.png"
    Image.new("RGB", (10, 10), color="red").save(image_one_path)
    doc.add_picture(str(image_one_path))

    doc.add_heading("Remote Control", level=1)
    doc.add_paragraph("The remote shall expose a dedicated input-select button.")
    doc.add_paragraph("See the second diagram below for button layout.")

    image_two_path = Path(tmp_dir) / "synthetic_blue.png"
    Image.new("RGB", (10, 10), color="blue").save(image_two_path)
    doc.add_picture(str(image_two_path))

out_path = Path(__file__).parent / "multi_figure.docx"
doc.save(str(out_path))
print(f"wrote {out_path}")
