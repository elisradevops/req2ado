"""Run once to (re)generate sample.pptx: `python build_pptx_fixture.py`."""
from pathlib import Path

from pptx import Presentation

presentation = Presentation()
layout = presentation.slide_layouts[1]  # "Title and Content"

slide1 = presentation.slides.add_slide(layout)
slide1.shapes.title.text = "Power Budget"
slide1.placeholders[1].text_frame.text = "The system shall not exceed 500W."

slide2 = presentation.slides.add_slide(layout)
slide2.shapes.title.text = "Thermal Requirements"
slide2.placeholders[1].text_frame.text = "The enclosure shall stay below 70C under full load."

out_path = Path(__file__).parent / "sample.pptx"
presentation.save(str(out_path))
print(f"wrote {out_path}")
