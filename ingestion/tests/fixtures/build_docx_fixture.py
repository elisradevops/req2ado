"""Run once to (re)generate sample.docx: `python build_docx_fixture.py`."""
from pathlib import Path

from docx import Document as DocxDocument

doc = DocxDocument()
doc.add_heading("Display", level=1)
doc.add_heading("Picture Modes", level=2)
doc.add_paragraph("The display shall support at least three picture presets.")

table = doc.add_table(rows=2, cols=2)
table.cell(0, 0).text = "Mode"
table.cell(0, 1).text = "Description"
table.cell(1, 0).text = "Vivid"
table.cell(1, 1).text = "High saturation preset"

doc.add_paragraph("See the diagram below for panel layout.")

out_path = Path(__file__).parent / "sample.docx"
doc.save(str(out_path))
print(f"wrote {out_path}")
