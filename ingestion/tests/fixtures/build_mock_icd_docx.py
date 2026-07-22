"""Run once to (re)generate mock_icd.docx: `python build_mock_icd_docx.py`.

Mock stand-in for the real engineering ICD reference document (unavailable
here — lives in a production environment). Embodies the real ICD convention
named in the original request: EACH top-level heading = Feature, everything
nested below it = Requirement. Uses real Word Heading styles (python-docx),
matching how the real ICD is actually authored, unlike the plain-numbered-
text style used for mock_tv_remote.docx.

Includes ONE genuine depth-3 sub-requirement ("Ripple Sub-Requirement" under
Power Distribution -> 3.3V Rail Requirements) to prove the docs/architecture.md
1g fix in `detect_profile_for_branch` (ingestion/app/leveling/autodetect.py):
a branch reaching depth 3 is still icd_style (Feature/Requirement/Requirement),
not misclassified as tv_style (which would have retyped the top two levels as
Epic/Epic). The Thermal Management branch is left at 2 levels, unchanged, so
this fixture still also proves the plain 2-level case.
"""
from pathlib import Path

from docx import Document as DocxDocument

doc = DocxDocument()

doc.add_heading("Power Distribution", level=1)
doc.add_heading("3.3V Rail Requirements", level=2)
doc.add_paragraph(
    "The 3.3V rail shall maintain output regulation within +/-2% under full "
    "rated load, verified per test procedure TP-114."
)
doc.add_heading("Ripple Sub-Requirement", level=3)
doc.add_paragraph(
    "Output ripple on the 3.3V rail shall not exceed 50mV peak-to-peak, "
    "measured per test procedure TP-114a."
)

doc.add_heading("Rail Sequencing Requirements", level=2)
doc.add_paragraph(
    "Rail sequencing shall enforce 3.3V rising before 1.2V to avoid latch-up "
    "on the FPGA power rail."
)

doc.add_heading("Thermal Management", level=1)
doc.add_heading("Fan Control", level=2)
doc.add_paragraph(
    "The system shall activate the chassis fan when the internal ambient "
    "temperature exceeds 45C, and deactivate it once the temperature falls "
    "below 40C."
)

out_path = Path(__file__).parent / "mock_icd.docx"
doc.save(str(out_path))
print(f"wrote {out_path}")
