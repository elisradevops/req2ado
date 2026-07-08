"""Run once to (re)generate mock_tv_remote.docx: `python build_mock_tv_remote_docx.py`.

Mock stand-in for the real TV-Receiver + Remote-Control worked example
(unavailable here — lives in a production environment). Every line is a
plain paragraph carrying its own numbering label in the text itself (e.g.
"1.1.1.1 <text>") rather than a Word Heading style — this is the real
"flat numbered list where one text line contains BOTH a Title and a
Description" convention named in the original request, and the harder of
the two real conventions (contrast with mock_icd.docx's Heading-style-based
document).

Two top-level branches on purpose, to exercise the SAME nominal depth
resolving to a DIFFERENT work-item type across branches of the SAME
document (the "Title 2 dual-role" case):
- Branch 1 ("1 ...") nests 5 levels deep (Epic/Epic/Feature/Requirement/
  sub-Requirement) -> classified tv_style (max depth >= 3).
- Branch 2 ("2 ...") only nests 2 levels deep -> classified icd_style
  (max depth <= 2), so its depth-1 node is a Feature, not an Epic — the
  same depth-1 position as branch 1's Epic.

Branch 1's depth-4 leaf is a deliberately long (>256 char) single line, to
exercise BASELINE title truncation (extract_baseline_title: truncate at a
word boundary, never summarize/paraphrase) on a real parsed document
rather than only a hand-built in-memory Block.
"""
from pathlib import Path

from docx import Document as DocxDocument

doc = DocxDocument()

# Branch 1 — deep, tv_style: Epic / sub-Epic / Feature / Requirement / sub-Requirement.
doc.add_paragraph("1 Remote Control")
doc.add_paragraph("1.1 Button Layout")
doc.add_paragraph("1.1.1 Volume Buttons")
doc.add_paragraph(
    "1.1.1.1 In accordance with the customer's stated operational expectations "
    "for user-facing volume control interactions across all supported input "
    "modes including infrared and Bluetooth HID pairing, the remote shall "
    "increase the currently active TV volume level by exactly one step within "
    "150 milliseconds of a single short button press."
)
doc.add_paragraph(
    "1.1.1.1.a The remote shall suppress repeat-trigger volume steps while the "
    "button is held continuously for longer than 500 milliseconds."
)

# Branch 2 — shallow, icd_style: same nominal depth-1 position as branch 1,
# but resolves to Feature (not Epic) because THIS branch never nests past
# depth 2.
doc.add_paragraph("2 Menu System")
doc.add_paragraph(
    "2.1 The on-screen menu shall list the five most recently used inputs "
    "under a Favorites section."
)

out_path = Path(__file__).parent / "mock_tv_remote.docx"
doc.save(str(out_path))
print(f"wrote {out_path}")
