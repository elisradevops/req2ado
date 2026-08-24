from pathlib import Path

import pytest
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_THEME_COLOR
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches

from app.canonical.model import BlockKind
from app.loaders.pptx_loader import load_pptx


def _build_pptx(tmp_path: Path) -> str:
    presentation = Presentation()
    title_layout = presentation.slide_layouts[1]  # "Title and Content"

    slide1 = presentation.slides.add_slide(title_layout)
    slide1.shapes.title.text = "Power Budget"
    body = slide1.placeholders[1]
    body.text_frame.text = "The system shall not exceed 500W."
    body.text_frame.add_paragraph().text = "The power supply shall support hot-swap."
    slide1.notes_slide.notes_text_frame.text = "Reviewed with the electrical lead."

    slide2 = presentation.slides.add_slide(title_layout)
    slide2.shapes.title.text = "Thermal Requirements"
    slide2.placeholders[1].text_frame.text = "The enclosure shall stay below 70C under full load."

    path = tmp_path / "sample.pptx"
    presentation.save(str(path))
    return str(path)


def test_load_pptx_produces_heading_per_slide_title(tmp_path: Path) -> None:
    doc = load_pptx(_build_pptx(tmp_path))
    headings = [b for b in doc.blocks if b.kind == BlockKind.HEADING]
    assert [h.text for h in headings] == ["Power Budget", "Thermal Requirements"]
    assert all(h.heading_level == 1 for h in headings)


def test_load_pptx_produces_paragraph_per_bullet(tmp_path: Path) -> None:
    doc = load_pptx(_build_pptx(tmp_path))
    paragraphs = [b for b in doc.blocks if b.kind == BlockKind.PARAGRAPH]
    texts = [p.text for p in paragraphs]
    assert "The system shall not exceed 500W." in texts
    assert "The power supply shall support hot-swap." in texts
    assert "The enclosure shall stay below 70C under full load." in texts


def test_load_pptx_produces_note_block_for_speaker_notes(tmp_path: Path) -> None:
    doc = load_pptx(_build_pptx(tmp_path))
    notes = [b for b in doc.blocks if b.kind == BlockKind.NOTE]
    assert len(notes) == 1
    assert notes[0].text == "Reviewed with the electrical lead."


def test_load_pptx_format_is_pptx(tmp_path: Path) -> None:
    doc = load_pptx(_build_pptx(tmp_path))
    assert doc.document.format.value == "pptx"


def test_load_pptx_blocks_are_in_slide_order(tmp_path: Path) -> None:
    doc = load_pptx(_build_pptx(tmp_path))
    offsets = [b.source_ref.start_offset for b in doc.blocks]
    assert offsets == sorted(offsets)


def test_load_pptx_sets_1_based_slide_number_as_page_for_every_block(tmp_path: Path) -> None:
    # Slide 1: title + 2 bullets + notes; slide 2: title + 1 bullet.
    doc = load_pptx(_build_pptx(tmp_path))
    slide1_blocks = doc.blocks[:4]
    slide2_blocks = doc.blocks[4:]
    assert all(b.source_ref.page == 1 for b in slide1_blocks)
    assert all(b.source_ref.page == 2 for b in slide2_blocks)


def test_load_pptx_captions_embedded_pictures(tmp_path: Path) -> None:
    from io import BytesIO

    from PIL import Image

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])  # blank layout
    image_buffer = BytesIO()
    Image.new("RGB", (10, 10), color="red").save(image_buffer, format="PNG")
    image_buffer.seek(0)
    slide.shapes.add_picture(image_buffer, Inches(1), Inches(1))
    path = tmp_path / "with_picture.pptx"
    presentation.save(str(path))

    calls = []

    def fake_describe_figure(image_bytes: bytes) -> str:
        calls.append(image_bytes)
        return "a red square"

    doc = load_pptx(str(path), describe_figure=fake_describe_figure)
    figure_blocks = [b for b in doc.blocks if b.kind == BlockKind.FIGURE]
    assert len(figure_blocks) == 1
    assert figure_blocks[0].text == "[FIGURE: a red square]"
    assert len(doc.figures) == 1
    assert doc.figures[0].caption == "a red square"
    assert len(calls) == 1
    # Single blank slide -> slide 1; a figure must carry the same
    # slide-number page as any text block would have on that slide.
    assert figure_blocks[0].source_ref.page == 1
    assert doc.figures[0].source_ref.page == 1


def test_load_pptx_skips_a_degenerate_2x2_image_without_calling_describe_figure(tmp_path: Path) -> None:
    # Mirrors pdf_loader.py's and docx_loader.py's own regression test for
    # the same real bug class: a pptx can embed a tiny rendering-trick
    # color-fill image the same way a PDF/docx can -- must never become a
    # Figure, never get a vision call, and never leave a gap in the
    # remaining figure ids.
    from io import BytesIO

    from PIL import Image

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])  # blank layout
    degenerate_buffer = BytesIO()
    Image.new("RGB", (2, 2), color="black").save(degenerate_buffer, format="PNG")
    degenerate_buffer.seek(0)
    slide.shapes.add_picture(degenerate_buffer, Inches(1), Inches(1))
    real_buffer = BytesIO()
    Image.new("RGB", (20, 20), color="red").save(real_buffer, format="PNG")
    real_buffer.seek(0)
    slide.shapes.add_picture(real_buffer, Inches(2), Inches(2))
    path = tmp_path / "mixed_figures.pptx"
    presentation.save(str(path))

    calls: list[bytes] = []

    def spy_describe_figure(image_bytes: bytes) -> str:
        calls.append(image_bytes)
        return "A real diagram."

    doc = load_pptx(str(path), describe_figure=spy_describe_figure)

    assert len(doc.figures) == 1  # the degenerate image contributed nothing
    assert doc.figures[0].caption == "A real diagram."
    assert len(calls) == 1  # never called on the 2x2 image
    figure_blocks = [b for b in doc.blocks if b.kind == BlockKind.FIGURE]
    assert len(figure_blocks) == 1


def test_load_pptx_dedupes_the_same_image_bytes_reused_across_slides(tmp_path: Path) -> None:
    # Gap C (mirrors pdf_loader.py's/docx_loader.py's own
    # seen_figure_hashes): a logo/watermark shape reused across slides has
    # no other filter to catch it in pptx -- the same image bytes must
    # collapse to one figure, not one per slide.
    from io import BytesIO

    from PIL import Image

    presentation = Presentation()
    logo_buffer = Image.new("RGB", (40, 40), color="black")
    real_buffer = Image.new("RGB", (40, 40), color="red")

    for pil_image in (logo_buffer, logo_buffer, real_buffer):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        buf = BytesIO()
        pil_image.save(buf, format="PNG")
        buf.seek(0)
        slide.shapes.add_picture(buf, Inches(1), Inches(1))

    path = tmp_path / "reused_image.pptx"
    presentation.save(str(path))

    calls: list[bytes] = []

    def spy_describe_figure(image_bytes: bytes) -> str:
        calls.append(image_bytes)
        return "A real diagram."

    doc = load_pptx(str(path), describe_figure=spy_describe_figure)

    assert len(doc.figures) == 2  # logo counted once, plus the distinct image
    assert len(calls) == 2  # the second logo occurrence never reaches vision
    figure_blocks = [b for b in doc.blocks if b.kind == BlockKind.FIGURE]
    assert len(figure_blocks) == 2


def _order_key(source_ref) -> tuple[int, int]:
    """Local stand-in for tree_builder._order_key (owned by other in-flight
    work, not imported here): document-order sort key of (page or 0,
    start_offset), used below to confirm pptx figures no longer cluster at
    the front of the whole deck via a hardcoded (0, 0) placeholder."""
    return (source_ref.page or 0, source_ref.start_offset)


def test_load_pptx_figure_orders_within_its_own_slide_not_at_deck_front(tmp_path: Path) -> None:
    from io import BytesIO

    from PIL import Image

    presentation = Presentation()
    title_layout = presentation.slide_layouts[1]

    slide1 = presentation.slides.add_slide(title_layout)
    slide1.shapes.title.text = "Power Budget"
    slide1.placeholders[1].text_frame.text = "The system shall not exceed 500W."

    slide2 = presentation.slides.add_slide(title_layout)
    slide2.shapes.title.text = "Thermal Requirements"
    slide2.placeholders[1].text_frame.text = "The enclosure shall stay below 70C under full load."
    image_buffer = BytesIO()
    Image.new("RGB", (10, 10), color="red").save(image_buffer, format="PNG")
    image_buffer.seek(0)
    slide2.shapes.add_picture(image_buffer, Inches(1), Inches(1))

    path = tmp_path / "mixed.pptx"
    presentation.save(str(path))

    doc = load_pptx(str(path))
    all_source_refs = [b.source_ref for b in doc.blocks] + [f.source_ref for f in doc.figures]
    sorted_refs = sorted(all_source_refs, key=_order_key)

    # The figure lives on slide 2, after slide 1's content -- it must NOT
    # sort to the front of the deck the way a hardcoded (0, 0)/page=None
    # placeholder would have forced it to.
    figure_ref = next(f.source_ref for f in doc.figures)
    assert sorted_refs.index(figure_ref) > 0
    assert [_order_key(r)[0] for r in sorted_refs] == sorted(_order_key(r)[0] for r in sorted_refs)


# ---------------------------------------------------------------------------
# Table extraction -- new capability (Phase 3 of
# docs/superpowers/specs/2026-08-12-table-fidelity-design.md): pptx
# previously extracted NO tables at all. Every builder below uses REAL
# python-pptx API calls (add_table/merge/fill.solid()/etc.) to build genuine
# fixtures, same "real fixture over hand-guessed shape" principle already
# used for docx/xlsx this session.
# ---------------------------------------------------------------------------


def _build_pptx_with_simple_table(tmp_path: Path) -> str:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])  # blank layout
    graphic_frame = slide.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(4), Inches(2))
    table = graphic_frame.table
    table.cell(0, 0).text_frame.text = "Name"
    table.cell(0, 1).text_frame.text = "Value"
    table.cell(1, 0).text_frame.text = "GPU"
    table.cell(1, 1).text_frame.text = "60"
    path = tmp_path / "with_table.pptx"
    presentation.save(str(path))
    return str(path)


def test_load_pptx_extracts_table_with_correct_text(tmp_path: Path) -> None:
    doc = load_pptx(_build_pptx_with_simple_table(tmp_path))
    assert len(doc.tables) == 1
    table = doc.tables[0]
    texts = [[cell.text if cell is not None else None for cell in row] for row in table.cells]
    assert texts == [["Name", "Value"], ["GPU", "60"]]


def test_load_pptx_table_source_ref_carries_slide_number_as_page(tmp_path: Path) -> None:
    doc = load_pptx(_build_pptx_with_simple_table(tmp_path))
    assert doc.tables[0].source_ref.page == 1


def test_load_pptx_horizontal_merge_produces_colspan_and_none_at_spanned_position(tmp_path: Path) -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    graphic_frame = slide.shapes.add_table(2, 3, Inches(1), Inches(1), Inches(6), Inches(2))
    table = graphic_frame.table
    # A real merge via python-pptx's own API -- table.cell(0,0).merge(...)
    # -- not a hand-guessed span shape.
    table.cell(0, 0).merge(table.cell(0, 1))
    for row_idx in range(2):
        for col_idx in range(3):
            cell = table.cell(row_idx, col_idx)
            if not cell.is_spanned:
                cell.text_frame.text = f"r{row_idx}c{col_idx}"
    path = tmp_path / "colspan.pptx"
    presentation.save(str(path))

    doc = load_pptx(str(path))
    row0 = doc.tables[0].cells[0]
    assert row0[0] is not None
    assert row0[0].text == "r0c0"
    assert row0[0].colspan == 2
    assert row0[0].rowspan == 1
    assert row0[1] is None  # spanned position -- absorbed by colspan, not its own cell
    assert row0[2] is not None and row0[2].colspan == 1


def test_load_pptx_vertical_merge_produces_rowspan_and_none_at_spanned_position(tmp_path: Path) -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    graphic_frame = slide.shapes.add_table(3, 2, Inches(1), Inches(1), Inches(4), Inches(3))
    table = graphic_frame.table
    table.cell(0, 0).merge(table.cell(1, 0))
    for row_idx in range(3):
        for col_idx in range(2):
            cell = table.cell(row_idx, col_idx)
            if not cell.is_spanned:
                cell.text_frame.text = f"r{row_idx}c{col_idx}"
    path = tmp_path / "rowspan.pptx"
    presentation.save(str(path))

    doc = load_pptx(str(path))
    cells = doc.tables[0].cells
    assert cells[0][0] is not None
    assert cells[0][0].rowspan == 2
    assert cells[0][0].colspan == 1
    assert cells[1][0] is None  # spanned position, vertically covered
    assert cells[2][0] is not None and cells[2][0].rowspan == 1


def test_load_pptx_is_header_true_for_first_row_when_table_first_row_flag_set(tmp_path: Path) -> None:
    # add_table's own default leaves table.first_row True (confirmed live
    # against the installed python-pptx) -- exercised here without an
    # explicit override to prove the DEFAULT deck-authoring case (a table
    # inserted via PowerPoint's own "Insert Table" UI) is handled correctly.
    doc = load_pptx(_build_pptx_with_simple_table(tmp_path))
    row0, row1 = doc.tables[0].cells
    assert all(cell is not None and cell.is_header for cell in row0)
    assert all(cell is not None and not cell.is_header for cell in row1)


def test_load_pptx_is_header_false_for_every_row_when_table_first_row_flag_unset(tmp_path: Path) -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    graphic_frame = slide.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(4), Inches(2))
    table = graphic_frame.table
    table.first_row = False
    table.cell(0, 0).text_frame.text = "a"
    table.cell(1, 0).text_frame.text = "b"
    path = tmp_path / "no_header.pptx"
    presentation.save(str(path))

    doc = load_pptx(str(path))
    assert all(cell is None or not cell.is_header for row in doc.tables[0].cells for cell in row)


def _build_pptx_with_bold_prefix_cell(tmp_path: Path) -> str:
    # Mirrors docx_loader's own bold-prefix regression fixture exactly: a
    # bold "Note:" prefix followed by non-bold body text in the SAME
    # paragraph must NOT bold the whole cell (uniform-bold rule), while a
    # wholly-bold cell with a trailing non-bold WHITESPACE-only run must
    # still read as bold.
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    graphic_frame = slide.shapes.add_table(1, 2, Inches(1), Inches(1), Inches(4), Inches(1))
    table = graphic_frame.table

    mixed_paragraph = table.cell(0, 0).text_frame.paragraphs[0]
    bold_prefix_run = mixed_paragraph.add_run()
    bold_prefix_run.text = "Note:"
    bold_prefix_run.font.bold = True
    rest_run = mixed_paragraph.add_run()
    rest_run.text = " rest of sentence"
    rest_run.font.bold = False

    bold_paragraph = table.cell(0, 1).text_frame.paragraphs[0]
    bold_run = bold_paragraph.add_run()
    bold_run.text = "AllBold"
    bold_run.font.bold = True
    trailing_whitespace_run = bold_paragraph.add_run()
    trailing_whitespace_run.text = "   "
    trailing_whitespace_run.font.bold = False

    path = tmp_path / "bold_prefix.pptx"
    presentation.save(str(path))
    return str(path)


def test_load_pptx_uniform_bold_rejects_bold_prefix_but_accepts_wholly_bold_cell(tmp_path: Path) -> None:
    doc = load_pptx(_build_pptx_with_bold_prefix_cell(tmp_path))
    row = doc.tables[0].cells[0]
    assert row[0] is not None and row[0].bold is False  # bold label prefix must NOT bold the whole cell
    assert row[1] is not None and row[1].bold is True  # wholly-bold text + trailing non-bold whitespace run


def test_load_pptx_alignment_mapping(tmp_path: Path) -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    graphic_frame = slide.shapes.add_table(1, 5, Inches(1), Inches(1), Inches(10), Inches(1))
    table = graphic_frame.table
    labels = ["unset", "left", "center", "right", "justify"]
    aligns = [None, PP_ALIGN.LEFT, PP_ALIGN.CENTER, PP_ALIGN.RIGHT, PP_ALIGN.JUSTIFY]
    for col_idx, (label, align) in enumerate(zip(labels, aligns)):
        cell = table.cell(0, col_idx)
        cell.text_frame.text = label
        if align is not None:
            cell.text_frame.paragraphs[0].alignment = align
    path = tmp_path / "alignment.pptx"
    presentation.save(str(path))

    doc = load_pptx(str(path))
    row = doc.tables[0].cells[0]
    assert [cell.align for cell in row] == [None, "left", "center", "right", "justify"]


def test_load_pptx_alignment_exotic_value_maps_to_none_not_guessed(tmp_path: Path) -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    graphic_frame = slide.shapes.add_table(1, 1, Inches(1), Inches(1), Inches(2), Inches(1))
    table = graphic_frame.table
    table.cell(0, 0).text_frame.text = "spread"
    table.cell(0, 0).text_frame.paragraphs[0].alignment = PP_ALIGN.DISTRIBUTE
    path = tmp_path / "distribute.pptx"
    presentation.save(str(path))

    doc = load_pptx(str(path))
    assert doc.tables[0].cells[0][0].align is None


def test_load_pptx_background_and_text_color_solid_rgb_extracted(tmp_path: Path) -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    graphic_frame = slide.shapes.add_table(1, 1, Inches(1), Inches(1), Inches(2), Inches(1))
    table = graphic_frame.table
    cell = table.cell(0, 0)
    cell.text_frame.text = "Colored"
    cell.fill.solid()
    cell.fill.fore_color.rgb = RGBColor(0x11, 0x22, 0x33)
    cell.text_frame.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    path = tmp_path / "colored.pptx"
    presentation.save(str(path))

    doc = load_pptx(str(path))
    extracted = doc.tables[0].cells[0][0]
    assert extracted.background_color == "#112233"
    # Unlike background, a genuine white TEXT color is a real, meaningful
    # choice and must NOT be normalized away.
    assert extracted.text_color == "#ffffff"


def test_load_pptx_white_background_normalizes_to_no_fill(tmp_path: Path) -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    graphic_frame = slide.shapes.add_table(1, 1, Inches(1), Inches(1), Inches(2), Inches(1))
    table = graphic_frame.table
    cell = table.cell(0, 0)
    cell.text_frame.text = "White"
    cell.fill.solid()
    cell.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    path = tmp_path / "white_fill.pptx"
    presentation.save(str(path))

    doc = load_pptx(str(path))
    assert doc.tables[0].cells[0][0].background_color is None


def test_load_pptx_theme_referenced_fill_and_font_color_degrade_to_none_not_raise(tmp_path: Path) -> None:
    # The two theme-color traps flagged in the design spec, both proven
    # live against the installed python-pptx before this test was written:
    # (1) a cell with NO explicit fill of its own (inherited from the
    # table's applied STYLE) has fill.type == None, and touching
    # .fore_color on it raises TypeError; (2) a cell with an explicit but
    # THEME-referenced fill/font color has fore_color.type/color.type ==
    # MSO_COLOR_TYPE.SCHEME, and touching .rgb on THAT raises
    # AttributeError. Neither must propagate out of load_pptx -- both must
    # degrade to background_color/text_color == None.
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    graphic_frame = slide.shapes.add_table(1, 2, Inches(1), Inches(1), Inches(4), Inches(1))
    table = graphic_frame.table

    inherited_cell = table.cell(0, 0)
    inherited_cell.text_frame.text = "Inherited"  # fill left untouched -- fill.type is None

    themed_cell = table.cell(0, 1)
    themed_cell.text_frame.text = "Themed"
    themed_cell.fill.solid()
    themed_cell.fill.fore_color.theme_color = MSO_THEME_COLOR.ACCENT_1
    themed_cell.text_frame.paragraphs[0].runs[0].font.color.theme_color = MSO_THEME_COLOR.ACCENT_2

    path = tmp_path / "theme_colors.pptx"
    presentation.save(str(path))

    doc = load_pptx(str(path))  # must not raise
    row = doc.tables[0].cells[0]
    assert row[0] is not None and row[0].background_color is None  # inherited from table style
    assert row[1] is not None and row[1].background_color is None  # theme-referenced fill
    assert row[1].text_color is None  # theme-referenced font color


def test_load_pptx_column_widths_are_fractions_summing_to_one(tmp_path: Path) -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    graphic_frame = slide.shapes.add_table(1, 3, Inches(1), Inches(1), Inches(6), Inches(1))
    table = graphic_frame.table
    table.columns[0].width = Inches(1)
    table.columns[1].width = Inches(2)
    table.columns[2].width = Inches(3)
    path = tmp_path / "widths.pptx"
    presentation.save(str(path))

    doc = load_pptx(str(path))
    widths = doc.tables[0].column_widths
    assert widths is not None
    assert widths == pytest.approx([1 / 6, 2 / 6, 3 / 6])
    assert sum(widths) == pytest.approx(1.0)


def test_load_pptx_border_style_always_none_no_reliable_per_cell_signal(tmp_path: Path) -> None:
    # pptx tables get their gridlines from an applied table-STYLE reference,
    # not explicit per-cell border XML -- there is no reliable signal to
    # read, so border_style must stay None (unknown), never a guessed
    # "grid". This is a deliberate, documented gap (see load_pptx's own
    # Table(...) construction comment), asserted explicitly here so it
    # reads as an intentional decision, not an untested oversight.
    doc = load_pptx(_build_pptx_with_simple_table(tmp_path))
    assert doc.tables[0].border_style is None


def test_load_pptx_table_style_provenance_is_none_not_heuristic(tmp_path: Path) -> None:
    # pptx table extraction is always explicit-markup-derived (or a
    # documented absence, like border_style) -- "heuristic" is reserved for
    # PDF's geometric/font-weight inference, never used here.
    doc = load_pptx(_build_pptx_with_simple_table(tmp_path))
    assert doc.tables[0].style_provenance is None
